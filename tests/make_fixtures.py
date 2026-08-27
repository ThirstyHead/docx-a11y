#!/usr/bin/env python3
"""Build the committed .docx test fixtures. Run once, then commit the outputs.

    .venv/bin/python tests/make_fixtures.py

Fixtures (tests/fixtures/):
  clean.docx          - fully compliant; expect 0 findings
  fixable.docx        - 6 fixable violations across all auto-fix rules
  violations.docx     - hard cases: heading skip, multiple H1, merged cell (manual)
  no-knead-bread.docx - real-world sample (copied from the original audit run)
"""
import io
import shutil
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

from docx_a11y.rules import _set_default_lang

FIX = Path(__file__).resolve().parent / "fixtures"
FIX.mkdir(exist_ok=True)


def _png_bytes() -> bytes:
    """Minimal 1x1 red PNG."""
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
    idat = chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
    iend = chunk(b"IEND", b"")
    return sig + ihdr + idat + iend


def _strip_default_lang(doc):
    """Remove w:lang from docDefaults (python-docx templates may carry one)."""
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        return
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        return
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        return
    lang = rpr.find(qn("w:lang"))
    if lang is not None:
        rpr.remove(lang)


def build_clean():
    doc = Document()
    h1 = doc.add_paragraph("Clean Document")
    h1.style = doc.styles["Heading 1"]
    doc.add_paragraph("Body paragraph with no issues.")
    h2 = doc.add_paragraph("Section")
    h2.style = doc.styles["Heading 2"]
    doc.add_paragraph("Another body paragraph.")
    _set_default_lang(doc.styles.element, "en-US")
    doc.core_properties.title = "Clean Document"
    doc.save(FIX / "clean.docx")


def build_fixable():
    """Layout (paragraph indexes):
       0 'Fixable Title'      Normal 20pt      -> map to Heading 1
       1 'Intro text ' + red run 'important part' (FF0000 on white = 4.0:1 < 4.5:1)
       2 image paragraph (no alt text)
       3 'Ingredients'        Normal           -> map to Heading 2
       4 'Instructions'       Normal           -> map to Heading 2
       then a 2x2 table with no tblHeader
       No dc:title, no docDefaults language.
    """
    doc = Document()
    p0 = doc.add_paragraph("Fixable Title")
    p0.runs[0].font.size = Pt(20)
    p1 = doc.add_paragraph("Intro text ")
    r = p1.add_run("important part")
    r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    doc.add_picture(io.BytesIO(_png_bytes()))  # its own paragraph, index 2
    doc.add_paragraph("Ingredients")
    doc.add_paragraph("Instructions")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "2"
    _strip_default_lang(doc)
    doc.core_properties.title = ""
    doc.save(FIX / "fixable.docx")


def build_violations():
    """Hard cases:
       0 'Violations Doc' Normal 20pt
       1 'Section A'  Heading 1
       2 'Deep'       Heading 3     <- level skip H1 -> H3 (fixable to H2)
       3 'Section B'  Heading 1     <- multiple H1 (manual)
       4 'Body'
       table 2x2, no tblHeader (fixable) + merged cell row0 (manual)
       No dc:title, no docDefaults language.
    """
    doc = Document()
    p0 = doc.add_paragraph("Violations Doc")
    p0.runs[0].font.size = Pt(20)
    h1a = doc.add_paragraph("Section A")
    h1a.style = doc.styles["Heading 1"]
    h3 = doc.add_paragraph("Deep")
    h3.style = doc.styles["Heading 3"]
    h1b = doc.add_paragraph("Section B")
    h1b.style = doc.styles["Heading 1"]
    doc.add_paragraph("Body")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "2"
    t.cell(0, 0).merge(t.cell(0, 1))  # gridSpan on row 0
    _strip_default_lang(doc)
    doc.core_properties.title = ""
    doc.save(FIX / "violations.docx")


def copy_bread():
    src = Path("/Users/scott/doc-remediation/No Knead Bread-test.docx")
    if src.exists():
        shutil.copy(src, FIX / "no-knead-bread.docx")
        print("copied no-knead-bread.docx")
    else:
        print(f"WARN: {src} not found; skipping no-knead-bread fixture")


if __name__ == "__main__":
    build_clean()
    build_fixable()
    build_violations()
    copy_bread()
    print("fixtures written to", FIX)