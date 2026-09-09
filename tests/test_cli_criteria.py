"""Test criteria configuration and what-if analysis via docx-a11y CLI."""
from pathlib import Path
from docx import Document
from docx.shared import RGBColor
import pytest
from docx_a11y.cli import main
from docx_a11y.audit import audit_file


def _create_low_contrast_doc(path: Path) -> Path:
    doc = Document()
    doc.core_properties.title = "Test Low Contrast Document"
    doc.add_heading("Test Heading 1", level=1)
    p = doc.add_paragraph()
    r = p.add_run("This is very low contrast light gray text.")
    r.font.color.rgb = RGBColor(220, 220, 220)  # Very low contrast on white
    doc.save(str(path))
    return path


def test_cli_init_criteria(tmp_path, monkeypatch):
    out_file = tmp_path / "criteria.txt"
    monkeypatch.setattr("sys.argv", ["docx-a11y", "--init-criteria", str(out_file)])
    with pytest.raises(SystemExit) as excinfo:
        main()
    assert excinfo.value.code == 0
    assert out_file.exists()
    content = out_file.read_text(encoding="utf-8")
    assert "[x] 1.4.3" in content
    assert "[x] 1.1.1" in content


def test_cli_criteria_what_if_exclusion(tmp_path, monkeypatch, capsys):
    doc_path = _create_low_contrast_doc(tmp_path / "low_contrast.docx")
    criteria_file = tmp_path / "criteria.txt"
    # Create checklist with 1.4.3 unselected for what-if scenario
    criteria_file.write_text(
        """# Custom Criteria Checklist
[x] 1.1.1 Non-text Content
[ ] 1.4.3 Contrast (Minimum)
[x] 2.4.2 Page Titled
""",
        encoding="utf-8",
    )

    out_dir = tmp_path / "reports"
    monkeypatch.setattr(
        "sys.argv",
        [
            "docx-a11y",
            str(doc_path),
            "--criteria",
            str(criteria_file),
            "--output-dir",
            str(out_dir),
            "--format",
            "md,json",
        ],
    )
    with pytest.raises(SystemExit) as excinfo:
        main()

    # Since the only violation was 1.4.3 and it was unchecked [ ],
    # it was routed to the excluded bucket, so active blocking is 0, exit code is 0 (pass)!
    assert excinfo.value.code == 0

    md_report = (out_dir / "low_contrast-a11y-report.md").read_text(encoding="utf-8")
    assert "What-If Analysis Active" in md_report
    assert "Excluded from compliance score via user criteria configuration" in md_report
    assert "1.4.3" in md_report
