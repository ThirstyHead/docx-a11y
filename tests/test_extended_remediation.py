"""Unit and integration tests for Phase 3 Extended Remediation Engine.

Tests:
1. Table Header Repeat and CantSplit remediation (w:tblHeader + w:cantSplit).
2. Deterministic color contrast scaling and preservation of color intent.
3. Heading structure normalization and visual pseudo-headings promotion.
4. Safe table unmerging utility for complex merged cells.
"""
from pathlib import Path
import docx
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from docx_a11y.audit import audit_file
from docx_a11y.contrast import adjust_color_for_contrast, contrast_ratio, hex_to_rgb, is_contrast_acceptable
from docx_a11y.rules import (
    AuditContext,
    ColorContrast,
    HeadingsNone,
    TableHeaderMissing,
    normalize_document_headings,
    unmerge_table_cells,
)

FIXTURES = Path(__file__).parent / "fixtures"


def test_contrast_adjustment_math():
    """Verify deterministic color contrast scaling."""
    # Red on white (initially ~4.0:1, below 4.5:1)
    adjusted_red = adjust_color_for_contrast("FF0000", "FFFFFF", target_ratio=4.5)
    fg_red = hex_to_rgb(adjusted_red)
    bg_white = hex_to_rgb("FFFFFF")
    assert fg_red is not None and bg_white is not None
    cr_red = contrast_ratio(fg_red, bg_white)
    assert cr_red >= 4.5
    assert is_contrast_acceptable(adjusted_red, "FFFFFF", is_large=False)

    # Light gray on white (initially ~1.2:1)
    adjusted_gray = adjust_color_for_contrast("E8E8E8", "FFFFFF", target_ratio=4.5)
    fg_gray = hex_to_rgb(adjusted_gray)
    assert fg_gray is not None
    cr_gray = contrast_ratio(fg_gray, bg_white)
    assert cr_gray >= 4.5

    # Dark blue on black (needs lightening)
    adjusted_dark = adjust_color_for_contrast("000033", "000000", target_ratio=4.5)
    fg_dark = hex_to_rgb(adjusted_dark)
    bg_black = hex_to_rgb("000000")
    assert fg_dark is not None and bg_black is not None
    cr_dark = contrast_ratio(fg_dark, bg_black)
    assert cr_dark >= 4.5


def test_table_header_and_cantsplit_remediation(tmp_path):
    """Verify TableHeaderMissing adds both w:tblHeader and w:cantSplit."""
    doc_path = tmp_path / "table_test.docx"
    doc = docx.Document()
    tbl = doc.add_table(rows=3, cols=2)
    tbl.rows[0].cells[0].text = "Header A"
    tbl.rows[0].cells[1].text = "Header B"
    tbl.rows[1].cells[0].text = "1"
    tbl.rows[1].cells[1].text = "2"
    tbl.rows[2].cells[0].text = "3"
    tbl.rows[2].cells[1].text = "4"
    doc.save(str(doc_path))

    # Audit initially flags table-header-missing
    res = audit_file(doc_path)
    finding = next(f for f in res["findings"] if f["rule_id"] == "table-header-missing")
    assert finding is not None

    # Apply fix
    rule = TableHeaderMissing()
    ctx = AuditContext()
    doc_to_fix = docx.Document(str(doc_path))
    success = rule.fix(doc_to_fix, type("F", (), finding)(), ctx)
    assert success is True

    fixed_path = tmp_path / "table_fixed.docx"
    doc_to_fix.save(str(fixed_path))

    # Inspect XML of first row
    reopened = docx.Document(str(fixed_path))
    tr0 = reopened.tables[0]._tbl.findall(qn("w:tr"))[0]
    trPr = tr0.find(qn("w:trPr"))
    assert trPr is not None
    assert trPr.find(qn("w:tblHeader")) is not None
    assert trPr.find(qn("w:cantSplit")) is not None


def test_color_contrast_remediation_adjusts_luminance(tmp_path):
    """Verify ColorContrast fix adjusts color value to pass contrast instead of discarding color."""
    doc_path = tmp_path / "contrast_test.docx"
    doc = docx.Document()
    p = doc.add_paragraph()
    r = p.add_run("Important warning text in bright red")
    r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)  # fails 4.5:1
    doc.save(str(doc_path))

    res = audit_file(doc_path)
    finding = next(f for f in res["findings"] if f["rule_id"] == "color-contrast")
    assert finding is not None

    rule = ColorContrast()
    ctx = AuditContext(background_rgb="FFFFFF")
    doc_to_fix = docx.Document(str(doc_path))
    success = rule.fix(doc_to_fix, type("F", (), finding)(), ctx)
    assert success is True

    fixed_path = tmp_path / "contrast_fixed.docx"
    doc_to_fix.save(str(fixed_path))

    # Re-audit: should pass with 0 contrast violations
    reaudit = audit_file(fixed_path)
    contrast_findings = [f for f in reaudit["findings"] if f["rule_id"] == "color-contrast"]
    assert len(contrast_findings) == 0

    # Inspect the run: color attribute should still exist, but scaled
    reopened = docx.Document(str(fixed_path))
    fixed_color = reopened.paragraphs[0].runs[0].font.color.rgb
    assert fixed_color is not None
    adjusted_hex = str(fixed_color)
    assert adjusted_hex != "FF0000"
    adj_rgb = hex_to_rgb(adjusted_hex)
    bg_rgb = hex_to_rgb("FFFFFF")
    assert adj_rgb is not None and bg_rgb is not None
    assert contrast_ratio(adj_rgb, bg_rgb) >= 4.5


def test_auto_promote_visual_pseudo_headings(tmp_path):
    """Verify auto_promote_headings promotes bold and large paragraphs to semantic Heading styles."""
    doc_path = tmp_path / "pseudo_headings.docx"
    doc = docx.Document()
    
    p1 = doc.add_paragraph()
    r1 = p1.add_run("Main Document Title")
    r1.font.size = Pt(20)
    r1.font.bold = True

    p2 = doc.add_paragraph("Regular introductory body text.")

    p3 = doc.add_paragraph()
    r3 = p3.add_run("Section One")
    r3.font.size = Pt(15)
    r3.font.bold = True

    p4 = doc.add_paragraph("Section body text.")
    doc.save(str(doc_path))

    # With default context, HeadingsNone.fix returns False
    rule = HeadingsNone()
    finding = type("F", (), {"rule_id": "headings-none", "location": "document body"})()
    
    ctx_default = AuditContext(auto_promote_headings=False)
    doc_default = docx.Document(str(doc_path))
    assert rule.fix(doc_default, finding, ctx_default) is False

    # With auto_promote_headings=True, pseudo-headings are promoted
    ctx_auto = AuditContext(auto_promote_headings=True)
    doc_auto = docx.Document(str(doc_path))
    success = rule.fix(doc_auto, finding, ctx_auto)
    assert success is True
    style0 = doc_auto.paragraphs[0].style
    style2 = doc_auto.paragraphs[2].style
    assert style0 is not None and style0.name == "Heading 1"
    assert style2 is not None and style2.name == "Heading 2"


def test_normalize_document_headings(tmp_path):
    """Verify heading normalization demotes multiple H1s and fills level skips."""
    doc = docx.Document()
    p1 = doc.add_paragraph("Title 1", style="Heading 1")
    p2 = doc.add_paragraph("Title 2", style="Heading 1")  # extra H1
    p3 = doc.add_paragraph("Sub-sub section", style="Heading 4")  # skip from H2 to H4

    count = normalize_document_headings(doc)
    assert count == 2
    s1, s2, s3 = p1.style, p2.style, p3.style
    assert s1 is not None and s1.name == "Heading 1"
    assert s2 is not None and s2.name == "Heading 2"
    assert s3 is not None and s3.name == "Heading 3"


def test_unmerge_table_cells(tmp_path):
    """Verify unmerge_table_cells removes gridSpan and vMerge XML nodes."""
    doc_path = FIXTURES / "violations.docx"
    doc = docx.Document(str(doc_path))
    tbl = doc.tables[0]

    # Count initial merged cell findings
    initial_res = audit_file(doc_path)
    merged_findings = [f for f in initial_res["findings"] if f["rule_id"] == "merged-cell"]
    assert len(merged_findings) > 0

    modified = unmerge_table_cells(tbl)
    assert modified > 0

    fixed_path = tmp_path / "unmerged.docx"
    doc.save(str(fixed_path))

    reaudit = audit_file(fixed_path)
    residual_merged = [f for f in reaudit["findings"] if f["rule_id"] == "merged-cell"]
    assert len(residual_merged) == 0
