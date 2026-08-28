"""Remediation engine: apply deterministic fixes from an audit result.

Safety model:
  - Works on a fresh in-memory Document copy of the source; never mutates the source file.
  - Only fixes flagged fixable=true are attempted, in deterministic order
    (structure before metadata, so the title rule can pick up a new H1).
  - Every fix has a precondition; failures are recorded, never exceptions.
  - Output is a new file; the original is untouched.
"""
import time
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from .audit import audit_file, load_result
from .findings import Finding
from .rules import RULES_BY_ID, RULES, AuditContext

# Deterministic application order (structure first, metadata last).
APPLY_ORDER = [
    "HeadingsNone",          # applies heading map if provided
    "HeadingSkipped",
    "ImageAltMissing",
    "TableHeaderMissing",
    "ColorContrast",
    "TitleMissing",          # after headings so dc:title can use the first H1
    "LanguageMissing",
]


@dataclass
class RemediationResult:
    output_path: str
    applied: list = field(default_factory=list)    # [rule_id, location]
    skipped: list = field(default_factory=list)    # [rule_id, location, reason]

    @property
    def ok(self):
        return not self.skipped

    def to_dict(self):
        return {"output_path": self.output_path, "applied": self.applied,
                "skipped": self.skipped, "ok": self.ok}


def apply_fixes(doc, result: dict, rr: RemediationResult, ctx: AuditContext) -> None:
    """Apply deterministic fixes for `result`'s fixable findings onto an in-memory
    Document. Populates rr.applied / rr.skipped. Never raises for rule failures."""
    findings = [Finding(**f) for f in result["findings"] if f.get("fixable")]
    # group by rule class name (derived from rule_id)
    by_rule = {}
    for f in findings:
        rcls = _class_for(f.rule_id)
        by_rule.setdefault(rcls, []).append(f)

    for rcls in APPLY_ORDER:
        rule = _rule_for_class(rcls)
        if rule is None:
            continue
        rule_findings = by_rule.get(rcls, [])
        # Rules whose fix() handles ALL of its findings in one shot (whole-doc mutations).
        bulk = rcls in ("HeadingsNone", "TitleMissing")
        if bulk and rule_findings:
            f0 = rule_findings[0]
            reason = None
            try:
                ok = rule.fix(doc, f0, ctx)
            except Exception as exc:
                ok = False
                reason = f"{type(exc).__name__}: {exc}"
            if ok:
                rr.applied.extend([rcls, f.location] for f in rule_findings)
            else:
                rr.skipped.extend([rcls, f.location, reason or "fix returned False"]
                                   for f in rule_findings)
            continue
        for f in rule_findings:
            reason = None
            try:
                ok = rule.fix(doc, f, ctx)
            except Exception as exc:
                ok = False
                reason = f"{type(exc).__name__}: {exc}"
            if ok:
                rr.applied.append([rcls, f.location])
            else:
                rr.skipped.append([rcls, f.location, reason or "fix returned False"])


def remediate(src_path, result_path, out_path, ctx=None) -> RemediationResult:
    src, out = Path(src_path), Path(out_path)
    res = load_result(result_path)
    doc = Document(str(src))
    if ctx is None:
        ctx = AuditContext(source_name=src.name)
    if not ctx.source_name:
        ctx.source_name = src.name

    rr = RemediationResult(output_path=str(out))
    apply_fixes(doc, res, rr, ctx)
    doc.save(str(out))
    return rr


def remediate_from_result(src_path, result: dict, out_path, ctx=None) -> RemediationResult:
    """Remediate using an in-memory audit result dict (no findings JSON file).

    Same safety model as remediate(): fresh in-memory Document copy, source
    untouched, output saved to out_path.
    """
    src, out = Path(src_path), Path(out_path)
    doc = Document(str(src))
    if ctx is None:
        ctx = AuditContext(source_name=src.name)
    if not ctx.source_name:
        ctx.source_name = src.name
    rr = RemediationResult(output_path=str(out))
    apply_fixes(doc, result, rr, ctx)
    doc.save(str(out))
    return rr


def fix_one(src_path, out_path=None, ctx=None) -> dict:
    """Audit -> remediate (if findings) -> verify, for one .docx.

    Never mutates src. Returns a JSON-safe dict:
      status: "pass" | "fail" | "error"
        pass  = re-audit has zero blocking findings (doc may still have
                non-blocking manual findings)
        fail  = re-audit still has blocking findings after fixes
        error = could not audit or save (corrupt/missing file)
      findings_before: int (findings on the source audit; 0 for clean docs)
      remediation: RemediationResult.to_dict() or None (None when 0 findings)
      reaudit: full re-audit result dict (None on error)
      error: str (error status only)
    """
    src = Path(src_path)
    if out_path is None:
        out_path = src.with_name(src.name + ".fixed.docx")
    out_path = Path(out_path)
    if ctx is None:
        ctx = AuditContext(source_name=src.name)
    if not ctx.source_name:
        ctx.source_name = src.name

    base = {"file": src.name, "output_path": str(out_path),
            "findings_before": 0, "remediation": None, "reaudit": None,
            "error": None}
    try:
        before = audit_file(src, ctx)
    except Exception as exc:
        return {**base, "status": "error",
                "error": f"audit failed: {type(exc).__name__}: {exc}"}
    base["findings_before"] = before["summary"]["total"]

    try:
        if before["findings"]:
            rr = remediate_from_result(src, before, out_path, ctx)
            base["remediation"] = rr.to_dict()
        else:
            # nothing to fix: still emit a copy so --out is always produced
            Document(str(src)).save(str(out_path))
        after = audit_file(out_path)
    except Exception as exc:
        return {**base, "status": "error",
                "error": f"remediate/verify failed: {type(exc).__name__}: {exc}"}

    base["reaudit"] = after
    base["status"] = "pass" if after["summary"]["pass"] else "fail"
    return base


def _batch_docx_files(directory: Path) -> list:
    """Sorted .docx files, non-recursive; skip Office lock files and our own .fixed.docx outputs."""
    if not directory.is_dir():
        raise NotADirectoryError(directory)
    files = [p for p in directory.iterdir()
             if p.is_file()
             and p.suffix.lower() == ".docx"
             and "~$" not in p.name
             and not p.name.endswith(".fixed.docx")]
    return sorted(files, key=lambda p: p.name)


def fix_batch(directory, ctx=None) -> dict:
    """fix_one() over every .docx in `directory` (non-recursive), continuing past
    per-file errors. Same ctx (e.g. heading_map, language) applies to all files.

    Returns JSON-safe dict:
      directory, started_at,
      entries: [ fix_one result dicts, in filename order ],
      summary: {total, pass, fail, error,
                findings_before (sum), findings_after (sum)}
    """
    d = Path(directory)
    files = _batch_docx_files(d)
    entries = []
    for p in files:
        # fresh ctx per file so source_name is correct; keep caller knobs
        fctx = AuditContext(source_name=p.name,
                            default_language=ctx.default_language if ctx else "en-US",
                            background_rgb=ctx.background_rgb if ctx else "FFFFFF",
                            heading_map=dict(ctx.heading_map) if ctx else {})
        entries.append(fix_one(p, None, fctx))
    s = {"total": len(entries),
         "pass": sum(1 for e in entries if e["status"] == "pass"),
         "fail": sum(1 for e in entries if e["status"] == "fail"),
         "error": sum(1 for e in entries if e["status"] == "error"),
         "findings_before": sum(e["findings_before"] for e in entries),
         "findings_after": sum(e["reaudit"]["summary"]["total"] for e in entries if e["reaudit"])}
    return {"directory": str(d),
            "started_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "entries": entries, "summary": s}


def _class_for(rule_id):
    # rule_ids: title-missing, language-missing, heading-level-skipped, headings-none,
    #           multiple-h1, image-alt-missing, table-header-missing, merged-cell, color-contrast
    mapping = {
        "title-missing": "TitleMissing",
        "language-missing": "LanguageMissing",
        "heading-level-skipped": "HeadingSkipped",
        "headings-none": "HeadingsNone",
        "multiple-h1": "MultipleH1",
        "image-alt-missing": "ImageAltMissing",
        "table-header-missing": "TableHeaderMissing",
        "merged-cell": "MergedCell",
        "color-contrast": "ColorContrast",
    }
    return mapping.get(rule_id, rule_id)


def _rule_for_class(rcls):
    for r in RULES:
        if type(r).__name__ == rcls:
            return r
    return None