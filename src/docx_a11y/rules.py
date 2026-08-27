"""Audit rules for WCAG 2.1 AA on .docx files.

Each Rule implements:
    check(doc, ctx) -> list[Finding]
    fix(doc, finding, ctx) -> bool     # deterministic fixes only; False = manual

Rules are ordered in RULES. Severity defaults can be overridden per rule.
Scope note (per local a11y wiki, word-a11y-workflow): keyboard/navigation
criteria (2.x) do not apply to .docx content; 2.4.2 (title) is the web SC
closest to document metadata naming and is applied by convention.
"""
import re
from dataclasses import dataclass, field
from typing import List, Optional

from docx.oxml.ns import qn

from .contrast import contrast_ratio, hex_to_rgb
from .findings import Finding

HEADING_RE = re.compile(r"^Heading (\d)$")
LANG_RE = re.compile(r"^[A-Za-z]{2,3}(-[A-Za-z0-9]{2,8})?$")
ALT_PLACEHOLDER = (
    "[ALT-NOT-PROVIDED: {kind} at {loc} in '{title}' - insert a human-written description]"
)


@dataclass
class AuditContext:
    """Knobs that make audits deterministic and caller-controlled."""

    source_name: str = "document.docx"
    default_language: str = "en-US"
    background_rgb: str = "FFFFFF"  # assumed page background for contrast math
    large_text_size_pt: float = 18.0
    large_text_bold_pt: float = 14.0
    # explicit heading map for deterministic structure assignment:
    # {paragraph_index: "Heading 1"} — overrides heuristics entirely when set.
    heading_map: dict = field(default_factory=dict)


def _para_text(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def _run_color_rgb(run) -> Optional[str]:
    c = run.font.color
    if c is not None and c.rgb is not None:
        return str(c.rgb)
    return None


def _effective_font_size_pt(run) -> Optional[float]:
    # explicit run size wins; None -> inherit (treated as body text 11pt for threshold)
    if run.font.size is not None:
        return run.font.size.pt
    return 11.0


def _is_large_text(size_pt: float, bold: bool) -> bool:
    return size_pt >= 18.0 or (bold and size_pt >= 14.0)


def _drawing_alt(drawing_el):
    """Return (has_alt, descr, title) for a w:drawing element."""
    for docPr in drawing_el.iter(qn("wp:docPr")):
        descr = (docPr.get("descr") or "").strip()
        title = (docPr.get("title") or "").strip()
        return (bool(descr or title), descr, title)
    return (False, "", "")


def _iter_images(doc):
    """Yield (kind, location, drawing_el) for every inline image."""
    n = 0
    for i, p in enumerate(doc.paragraphs):
        for d in p._p.iter(qn("w:drawing")):
            n += 1
            yield ("body image", f"paragraph[{i}] drawing#{n}", d)
    for si, sec in enumerate(doc.sections):
        for name, hf in (("header", sec.header), ("footer", sec.footer)):
            for j, p in enumerate(hf.paragraphs):
                for d in p._p.iter(qn("w:drawing")):
                    n += 1
                    yield (f"{name} image", f"section[{si}] {name} paragraph[{j}] drawing#{n}", d)


def _default_lang(styles_el):
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        return None
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        return None
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        return None
    lang = rpr.find(qn("w:lang"))
    return lang.get(qn("w:val")) if lang is not None else None


def _set_default_lang(styles_el, val):
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        from docx.oxml import OxmlElement
        dd = OxmlElement("w:docDefaults")
        styles_el.insert(0, dd)
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        from docx.oxml import OxmlElement
        rprd = OxmlElement("w:rPrDefault")
        dd.append(rprd)
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        from docx.oxml import OxmlElement
        rpr = OxmlElement("w:rPr")
        rprd.append(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        from docx.oxml import OxmlElement
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), val)


# ---------------------------------------------------------------------------
# Rules
# ---------------------------------------------------------------------------

class TitleMissing:
    """SC 2.4.2 (by convention): document title metadata must be non-empty."""

    rule_id = "title-missing"
    sc = "2.4.2"
    severity = "moderate"

    def check(self, doc, ctx):
        title = (doc.core_properties.title or "").strip()
        if title:
            return []
        return [Finding(self.rule_id, self.sc, self.severity, "docProps/core.xml",
                        "Document has no title set (dc:title empty).",
                        "cp.title = ''", True,
                        "Set dc:title from first Heading 1 or filename stem.")]

    def fix(self, doc, finding, ctx):
        candidates = []
        for p in doc.paragraphs:
            if p.style is not None and p.style.name == "Heading 1":
                t = _para_text(p).strip()
                if t:
                    candidates.append(t)
                    break
        if not candidates:
            candidates.append(ctx.source_name.rsplit(".", 1)[0].replace("-", " ").replace("_", " "))
        doc.core_properties.title = candidates[0]
        return True


class LanguageMissing:
    """SC 3.1.1: default language must be declared in docDefaults."""

    rule_id = "language-missing"
    sc = "3.1.1"
    severity = "moderate"

    def check(self, doc, ctx):
        val = _default_lang(doc.styles.element)
        if val and LANG_RE.fullmatch(val):
            return []
        return [Finding(self.rule_id, self.sc, self.severity, "word/styles.xml w:docDefaults",
                        "No valid default language declared (w:docDefaults w:lang missing or malformed).",
                        f"w:lang val={val!r}", True,
                        f"Set w:lang w:val='{ctx.default_language}' on docDefaults rPr.")]

    def fix(self, doc, finding, ctx):
        _set_default_lang(doc.styles.element, ctx.default_language)
        return True


class HeadingSkipped:
    """SC 1.3.1: heading level must not skip (e.g. H1 -> H3)."""

    rule_id = "heading-level-skipped"
    sc = "1.3.1"
    severity = "serious"

    def check(self, doc, ctx):
        out, prev = [], None
        for i, p in enumerate(doc.paragraphs):
            style = p.style.name if p.style is not None else ""
            m = HEADING_RE.fullmatch(style)
            if not m:
                continue
            lvl = int(m.group(1))
            if prev is not None and lvl - prev > 1:
                out.append(Finding(self.rule_id, self.sc, self.severity,
                                   f"paragraph[{i}]",
                                   f"Heading level skipped: H{prev} -> H{lvl}.",
                                   f"style='{style}', text={_para_text(p)[:60]!r}", True,
                                   f"Reassign to 'Heading {prev + 1}' (or lower subsequent headings)."))
            prev = lvl
        return out

    def fix(self, doc, finding, ctx):
        i = int(finding.location.split("[")[1].split("]")[0])
        p = doc.paragraphs[i]
        prev = None
        for q in doc.paragraphs[:i]:
            style = q.style.name if q.style is not None else ""
            m = HEADING_RE.fullmatch(style)
            if m:
                prev = int(m.group(1))
        if prev is None:
            return False
        p.style = doc.styles[f"Heading {prev + 1}"]
        return True


class HeadingsNone:
    """SC 1.3.1: document should use heading styles for structure."""

    rule_id = "headings-none"
    sc = "1.3.1"
    severity = "serious"

    def check(self, doc, ctx):
        # Always flag when no heading styles exist, even if ctx.heading_map is
        # set: the map is the *remedy* (fix applies it), not an excuse to hide
        # the finding. One-shot workflows (fix_one) audit and remediate with
        # the same ctx, so suppressing here would starve the fix of its finding.
        any_heading = any(
            p.style is not None and HEADING_RE.fullmatch(p.style.name or "")
            for p in doc.paragraphs
        )
        if any_heading:
            return []
        return [Finding(self.rule_id, self.sc, self.severity, "document body",
                        "No Heading styles used in the document; structure is presentation-only.",
                        "0 paragraphs with style 'Heading N'", True,
                        "Provide --heading-map '<idx>=<Heading N>' entries (deterministic) "
                        "or apply native heading styles in Word.")]

    def fix(self, doc, finding, ctx):
        if not ctx.heading_map:
            return False
        for idx_str, style in sorted(ctx.heading_map.items(), key=lambda kv: int(kv[0])):
            idx = int(idx_str)
            doc.paragraphs[idx].style = doc.styles[style]
        return True


class MultipleH1:
    """SC 1.3.1 (best practice): single H1 for the document title."""

    rule_id = "multiple-h1"
    sc = "1.3.1"
    severity = "moderate"

    def check(self, doc, ctx):
        h1 = [i for i, p in enumerate(doc.paragraphs)
              if p.style is not None and p.style.name == "Heading 1"]
        if len(h1) > 1:
            return [Finding(self.rule_id, self.sc, self.severity,
                            f"paragraph[{h1[0]}] and paragraph[{h1[1]}] (+{len(h1) - 2} more)",
                            f"{len(h1)} Heading 1 paragraphs; recommend a single H1.",
                            f"indexes={h1}", False,
                            "Keep one H1 for the document title; demote the rest to H2.")]
        return []

    def fix(self, doc, finding, ctx):
        return False


class ImageAltMissing:
    """SC 1.1.1: every inline image needs alt text (wp:docPr descr/title)."""

    rule_id = "image-alt-missing"
    sc = "1.1.1"
    severity = "critical"

    def check(self, doc, ctx):
        out = []
        for kind, loc, d in _iter_images(doc):
            has_alt, _, _ = _drawing_alt(d)
            if not has_alt:
                out.append(Finding(self.rule_id, self.sc, self.severity, loc,
                                   f"{kind} has no alternative text (wp:docPr descr/title empty).",
                                   "wp:docPr descr='' title=''", True,
                                   "Write a human description; auto-fix inserts an "
                                   "[ALT-NOT-PROVIDED: ...] marker to keep the location tracked."))
        return out

    def fix(self, doc, finding, ctx):
        # locate the drawing by re-iterating in the same order
        target_loc = finding.location
        for kind, loc, d in _iter_images(doc):
            if loc != target_loc:
                continue
            for docPr in d.iter(qn("wp:docPr")):
                title = (doc.core_properties.title or ctx.source_name).strip()
                docPr.set("descr", ALT_PLACEHOLDER.format(kind=kind, loc=loc, title=title))
                if not (docPr.get("title") or "").strip():
                    docPr.set("title", ALT_PLACEHOLDER.format(kind=kind, loc=loc, title=title))
            return True
        return False


class TableHeaderMissing:
    """SC 1.3.1: data tables need a designated header row (w:tblHeader)."""

    rule_id = "table-header-missing"
    sc = "1.3.1"
    severity = "serious"

    def check(self, doc, ctx):
        out = []
        for ti, table in enumerate(doc.tables):
            tbl_el = table._tbl
            rows = tbl_el.findall(qn("w:tr"))
            if not rows:
                continue
            header_rows = 0
            for tr in rows:
                trPr = tr.find(qn("w:trPr"))
                if trPr is not None and trPr.find(qn("w:tblHeader")) is not None:
                    header_rows += 1
            if header_rows == 0:
                out.append(Finding(self.rule_id, self.sc, self.severity,
                                   f"table[{ti}]",
                                   f"Table has no designated header row (w:tblHeader absent).",
                                   f"rows={len(rows)}, tblHeader rows=0", True,
                                   "Mark first row as header (w:tblHeader) and ensure each "
                                   "column has a header cell."))
        return out

    def fix(self, doc, finding, ctx):
        ti = int(finding.location.split("[")[1].split("]")[0])
        tbl_el = doc.tables[ti]._tbl
        first_tr = tbl_el.findall(qn("w:tr"))[0]
        trPr = first_tr.find(qn("w:trPr"))
        if trPr is None:
            from docx.oxml import OxmlElement
            trPr = OxmlElement("w:trPr")
            first_tr.insert(0, trPr)
        if trPr.find(qn("w:tblHeader")) is None:
            from docx.oxml import OxmlElement
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trPr.append(th)
        return True


class MergedCell:
    """SC 1.3.1: merged cells (gridSpan/vMerge) may break table semantics."""

    rule_id = "merged-cell"
    sc = "1.3.1"
    severity = "moderate"

    def check(self, doc, ctx):
        out = []
        for ti, table in enumerate(doc.tables):
            for ri, tr in enumerate(table._tbl.findall(qn("w:tr"))):
                for ci, tc in enumerate(tr.findall(qn("w:tc"))):
                    tcPr = tc.find(qn("w:tcPr"))
                    if tcPr is None:
                        continue
                    if tcPr.find(qn("w:gridSpan")) is not None or tcPr.find(qn("w:vMerge")) is not None:
                        out.append(Finding(self.rule_id, self.sc, self.severity,
                                           f"table[{ti}] row[{ri}] cell[{ci}]",
                                           "Merged cell (gridSpan/vMerge) may break screen-reader table structure.",
                                           f"gridSpan/vMerge present", False,
                                           "Unmerge and repeat values unless the data truly spans."))
        return out

    def fix(self, doc, finding, ctx):
        return False


class ColorContrast:
    """SC 1.4.3: hardcoded run colors must meet 4.5:1 (3:1 large text) vs background."""

    rule_id = "color-contrast"
    sc = "1.4.3"
    severity = "moderate"

    def check(self, doc, ctx):
        out = []
        bg = hex_to_rgb(ctx.background_rgb)
        for i, p in enumerate(doc.paragraphs):
            for r in p.runs:
                rgb = _run_color_rgb(r)
                if rgb is None or not r.text.strip():
                    continue
                fg = hex_to_rgb(rgb)
                if fg is None or bg is None:
                    continue
                ratio = contrast_ratio(fg, bg)
                size = _effective_font_size_pt(r)
                large = _is_large_text(size, bool(r.font.bold))
                threshold = 3.0 if large else 4.5
                if ratio < threshold:
                    out.append(Finding(self.rule_id, self.sc, self.severity,
                                       f"paragraph[{i}]",
                                       f"Text color #{rgb} fails contrast: {ratio:.2f}:1 < {threshold}:1 "
                                       f"(assumed background #{ctx.background_rgb}).",
                                       f"run={r.text[:40]!r}, color={rgb}, size={size}pt, "
                                       f"bold={bool(r.font.bold)}, large={large}", True,
                                       "Remove the w:color override (text inherits default color) "
                                       "or choose a color meeting the threshold."))
        return out

    def fix(self, doc, finding, ctx):
        i = int(finding.location.split("[")[1].split("]")[0])
        p = doc.paragraphs[i]
        fixed = 0
        for r in p.runs:
            rgb = _run_color_rgb(r)
            if rgb is None:
                continue
            fg = hex_to_rgb(rgb)
            bg = hex_to_rgb(ctx.background_rgb)
            if fg is None or bg is None:
                continue
            if contrast_ratio(fg, bg) < (3.0 if _is_large_text(_effective_font_size_pt(r), bool(r.font.bold)) else 4.5):
                rpr = r._r.find(qn("w:rPr"))
                col = rpr.find(qn("w:color")) if rpr is not None else None
                if col is not None:
                    rpr.remove(col)
                    fixed += 1
        return fixed > 0


RULES = [
    TitleMissing(),
    LanguageMissing(),
    HeadingSkipped(),
    HeadingsNone(),
    MultipleH1(),
    ImageAltMissing(),
    TableHeaderMissing(),
    MergedCell(),
    ColorContrast(),
]

RULES_BY_ID = {r.rule_id: r for r in RULES}